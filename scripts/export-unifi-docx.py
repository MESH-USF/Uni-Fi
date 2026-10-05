#!/usr/bin/env python3
"""Convert a Marked lexer AST and rendered Mermaid PNGs to a Word report.

Usage:
  python scripts/export-unifi-docx.py --ast /tmp/unifi-document-ast.json \
    --figures output/playwright/unifi-document --output docs/Uni-Fi.docx

Requires python-docx. The AST must be the JSON serialization of
marked.lexer(markdown); diagrams are figure-01.png, figure-02.png, etc.
No Markdown prose, table cells, code examples, or list items are summarized.
Mermaid source is replaced with its rendered image, in the original order.
"""

from __future__ import annotations

import argparse
import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run


FIGURE_TITLES = (
    "The standalone radios and their optional companion clients",
    "The three protocol layers and the end-to-end data path",
    "Minimal firmware build selection",
    "Firmware module responsibilities and connections",
    "Device startup and provisioning",
    "Browser connection and companion handshake",
    "Incoming-message processing and client collection",
    "The complete SOS and human receipt exchange",
    "What each incident state does and does not prove",
    "Planned GPS and map data ownership",
)


def element(tag: str, **attributes: object):
    result = OxmlElement(tag)
    for key, value in attributes.items():
        result.set(qn(key), str(value))
    return result


def slug(text: str) -> str:
    # GitHub-compatible anchors for the headings in this report.
    return re.sub(r"[^\w\- ]", "", text.lower()).replace(" ", "-")


def tokens_text(tokens: list[dict]) -> str:
    result = []
    for token in tokens:
        if token.get("tokens"):
            result.append(tokens_text(token["tokens"]))
        else:
            result.append(html.unescape(str(token.get("text", ""))))
    return "".join(result)


class Exporter:
    def __init__(self, tokens: list[dict], figures: Path):
        self.doc = Document()
        self.tokens = tokens
        self.figures = figures
        self.figure_count = 0
        self.block_count = 0
        self.bookmark_count = 0
        self.heading = ""
        self.anchors = {}
        for token in tokens:
            if token["type"] == "heading":
                text = token.get("text", tokens_text(token.get("tokens", [])))
                self.anchors[slug(text)] = "section_" + str(len(self.anchors) + 1)
        self.configure()

    def configure(self):
        section = self.doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.685)
        section.right_margin = Inches(0.685)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.3)
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Liberation Sans"
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for level, size in ((1, 19), (2, 14), (3, 12), (4, 11)):
            style = styles[f"Heading {level}"]
            style.font.name = "Liberation Sans"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string("173E4D")
            style.paragraph_format.space_before = Pt(15 if level < 3 else 10)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        styles["Caption"].font.name = "Liberation Sans"
        styles["Caption"].font.size = Pt(9)
        styles["Caption"].font.color.rgb = RGBColor.from_string("355260")
        styles["Caption"].paragraph_format.space_after = Pt(10)
        for name in ("List Bullet", "List Number"):
            styles[name].paragraph_format.space_after = Pt(4)
        self.doc.core_properties.title = (
            "Uni-Fi: structure, implementation and protocol walkthrough"
        )
        self.doc.core_properties.subject = "Heltec V3 prototype technical report"
        self.doc.core_properties.author = "MeshUSF"
        self.doc.core_properties.keywords = "Uni-Fi, MeshCore, Heltec V3, LoRa"

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("UNI-FI  /  MESHUSF  •  TECHNICAL WALKTHROUGH")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("55717D")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("GPS-free prototype  |  Page ").font.size = Pt(8)
        self.field(footer, "PAGE")
        footer.add_run(" of ").font.size = Pt(8)
        self.field(footer, "NUMPAGES")
        section.different_first_page_header_footer = True
        self.doc.settings.element.append(element("w:updateFields", **{"w:val": "true"}))

    @staticmethod
    def field(paragraph, instruction: str):
        run = paragraph.add_run()
        run.font.size = Pt(8)
        run._r.append(element("w:fldChar", **{"w:fldCharType": "begin"}))
        text = element("w:instrText", **{"xml:space": "preserve"})
        text.text = f" {instruction} "
        run._r.append(text)
        run._r.append(element("w:fldChar", **{"w:fldCharType": "end"}))

    @staticmethod
    def add_run(parent, text: str):
        if hasattr(parent, "add_run"):
            return parent.add_run(text)
        r = element("w:r")
        parent.append(r)
        run = Run(r, None)
        run.text = text
        return run

    def inline(self, parent, tokens: list[dict], bold=False, italic=False, linked=False):
        for token in tokens:
            kind = token["type"]
            if kind in ("strong", "em", "del"):
                self.inline(
                    parent,
                    token.get("tokens", [{"type": "text", "text": token.get("text", "")}]),
                    bold=bold or kind == "strong",
                    italic=italic or kind == "em",
                    linked=linked,
                )
                continue
            if kind == "link":
                href = token["href"]
                link = element("w:hyperlink")
                if href.startswith("#"):
                    anchor = self.anchors.get(href[1:])
                    if not anchor:
                        raise ValueError(f"Unresolved document anchor: {href}")
                    link.set(qn("w:anchor"), anchor)
                    link.set(qn("w:history"), "1")
                else:
                    rel = self.doc.part.relate_to(
                        href,
                        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                        is_external=True,
                    )
                    link.set(qn("r:id"), rel)
                target = parent._p if hasattr(parent, "_p") else parent
                target.append(link)
                self.inline(link, token.get("tokens", [{"type": "text", "text": token.get("text", href)}]), bold, italic, True)
                continue
            if kind == "br":
                run = self.add_run(parent, "")
                run.add_break()
                continue
            if kind == "text" and token.get("tokens"):
                self.inline(parent, token["tokens"], bold, italic, linked)
                continue
            if kind not in ("text", "escape", "codespan", "html"):
                raise ValueError(f"Unsupported inline token: {kind}")
            # Marked HTML-escapes inline code as well as ordinary text.
            text = html.unescape(str(token.get("text", token.get("raw", ""))))
            text = text.replace(
                "Diagrams use Mermaid, which GitHub renders; directory trees\nand byte layouts are plain text.",
                "Charts are rendered in this document; the authoritative Markdown source uses Mermaid. Directory trees and byte layouts are plain text.",
            )
            if kind != "codespan":
                # Soft Markdown line breaks are normal spaces, unlike code blocks.
                text = re.sub(r"\s*\n\s*", " ", text)
            run = self.add_run(parent, text)
            run.bold = bold
            run.italic = italic
            if linked:
                run.font.color.rgb = RGBColor.from_string("137385")
                run.underline = True
            if kind == "codespan":
                run.font.name = "Liberation Mono"
                run.font.size = Pt(8.3)
                properties = run._r.get_or_add_rPr()
                properties.append(element("w:shd", **{"w:fill": "EEF3F5"}))

    def paragraph(self, token, style=None):
        paragraph = self.doc.add_paragraph(style=style)
        self.inline(paragraph, token.get("tokens", [{"type": "text", "text": token.get("text", "")}]))
        return paragraph

    def bookmark(self, paragraph, key):
        self.bookmark_count += 1
        paragraph._p.insert(0, element("w:bookmarkStart", **{"w:id": self.bookmark_count, "w:name": self.anchors[key]}))
        paragraph._p.append(element("w:bookmarkEnd", **{"w:id": self.bookmark_count}))

    def cover(self):
        self.doc.add_paragraph().paragraph_format.space_after = Pt(80)
        p = self.doc.add_paragraph()
        run = p.add_run("UNI-FI")
        run.bold = True
        run.font.size = Pt(42)
        run.font.color.rgb = RGBColor.from_string("173E4D")
        p = self.doc.add_paragraph("Mesh network developed by MeshUSF")
        p.runs[0].font.size = Pt(16)
        p.paragraph_format.space_after = Pt(28)
        p = self.doc.add_paragraph("Structure, implementation\nand protocol walkthrough")
        p.runs[0].font.size = Pt(24)
        p.paragraph_format.space_after = Pt(32)
        self.doc.add_paragraph("Technical report • Documentation date: 5 October 2026")
        self.doc.add_paragraph("Implementation snapshot: fd8a149 • feat/heltec-v3-prototype")
        p = self.doc.add_paragraph()
        p.add_run("CURRENT STATUS\n").bold = True
        p.add_run("GPS-free Heltec V3 prototype. Firmware builds and host/browser tests pass; physical RF, BLE, button, OLED and battery acceptance remain pending.")
        p.paragraph_format.space_before = Pt(32)
        self.doc.add_paragraph("This report preserves the complete technical walkthrough. Diagrams are rendered figures; the reading guide links to document sections. All example keys, names and IDs are illustrative, not deployment secrets.")
        self.doc.add_page_break()

    def code(self, token):
        if str(token.get("lang", "")).strip().split(" ")[0] == "mermaid":
            self.figure()
            return
        text = token.get("text", "")
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.widow_control = False
        paragraph.paragraph_format.keep_together = False
        paragraph._p.get_or_add_pPr().append(element("w:shd", **{"w:fill": "F1F5F7"}))
        run = paragraph.add_run(text)
        run.font.name = "Liberation Mono"
        longest = max((len(line) for line in text.splitlines()), default=0)
        run.font.size = Pt(8.1 if longest < 106 else 7.2)

    def figure(self):
        self.figure_count += 1
        path = self.figures / f"figure-{self.figure_count:02d}.png"
        if not path.is_file():
            raise FileNotFoundError(f"Missing rendered Mermaid diagram: {path}")
        image = Image.from_file(str(path))
        aspect = image.px_width / image.px_height
        landscape = aspect > 1.7
        if landscape:
            section = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Inches(11.69)
            section.page_height = Inches(8.27)
            section.different_first_page_header_footer = False
        width = min(10.32, 5.65 * aspect) if landscape else min(6.9, 8.55 * aspect)
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.page_break_before = False
        paragraph.paragraph_format.keep_with_next = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(width))
        title = (FIGURE_TITLES[self.figure_count - 1] if self.figure_count <= len(FIGURE_TITLES) else self.heading)
        caption = self.doc.add_paragraph(f"Figure {self.figure_count}. {title}", style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if landscape:
            section = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = Inches(8.27)
            section.page_height = Inches(11.69)
            section.different_first_page_header_footer = False

    def table(self, token):
        header = token["header"]
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        count = len(header)
        if count == 3:
            widths = (1.65, 2.25, 3.0)
        elif count == 4:
            widths = (1.0, 1.85, 0.9, 3.15)
        elif count == 2:
            widths = (2.0, 4.9)
        else:
            widths = (6.9 / count,) * count
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        table.rows[0]._tr.get_or_add_trPr().append(element("w:tblHeader"))
        all_rows = [header] + token.get("rows", [])
        for row_index, row_data in enumerate(all_rows):
            row = table.rows[0] if row_index == 0 else table.add_row()
            row._tr.get_or_add_trPr().append(element("w:cantSplit"))
            for index, cell_token in enumerate(row_data):
                cell = row.cells[index]
                cell.width = Inches(widths[index])
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.05
                self.inline(paragraph, cell_token.get("tokens", [{"type": "text", "text": cell_token.get("text", "")}]), bold=row_index == 0)
                for run in paragraph.runs:
                    if run.font.name != "Liberation Mono":
                        run.font.size = Pt(9)
                    else:
                        run.font.size = Pt(7.7)
                if row_index == 0:
                    cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "E2EDF0"}))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def numbering(self, start: int = 1):
        root = self.doc.part.numbering_part.element
        abstracts = root.findall(qn("w:abstractNum"))
        abstract_id = max((int(n.get(qn("w:abstractNumId"))) for n in abstracts), default=-1) + 1
        abstract = element("w:abstractNum", **{"w:abstractNumId": abstract_id})
        level = element("w:lvl", **{"w:ilvl": 0})
        level.append(element("w:start", **{"w:val": start}))
        level.append(element("w:numFmt", **{"w:val": "decimal"}))
        level.append(element("w:lvlText", **{"w:val": "%1."}))
        level.append(element("w:lvlJc", **{"w:val": "left"}))
        props = element("w:pPr")
        props.append(element("w:ind", **{"w:left": 360, "w:hanging": 240}))
        level.append(props)
        abstract.append(level)
        root.append(abstract)
        number_id = max((int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))), default=0) + 1
        number = element("w:num", **{"w:numId": number_id})
        number.append(element("w:abstractNumId", **{"w:val": abstract_id}))
        root.append(number)
        return number_id

    def list(self, token, depth=0):
        ordered = token.get("ordered", False)
        number_id = self.numbering(int(token.get("start") or 1)) if ordered else None
        for item in token.get("items", []):
            blocks = item.get("tokens", [{"type": "text", "text": item.get("text", "")}])
            first = True
            for block in blocks:
                if block["type"] == "space":
                    continue
                if block["type"] == "list":
                    self.list(block, depth + 1)
                    continue
                if block["type"] in ("text", "paragraph"):
                    paragraph = self.paragraph(block, style=("List Number" if ordered else "List Bullet") if first else None)
                    if depth:
                        paragraph.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                    if ordered and first:
                        props = paragraph._p.get_or_add_pPr()
                        num = element("w:numPr")
                        num.append(element("w:ilvl", **{"w:val": 0}))
                        num.append(element("w:numId", **{"w:val": number_id}))
                        props.append(num)
                    first = False
                else:
                    self.blocks([block])

    def blocks(self, tokens):
        for token in tokens:
            kind = token["type"]
            if kind == "space":
                continue
            self.block_count += 1
            if kind == "heading":
                self.heading = token.get("text", "")
                depth = token["depth"]
                paragraph = self.paragraph(token, style=f"Heading {min(depth, 4)}")
                self.bookmark(paragraph, slug(self.heading))
                if depth == 2 and self.heading != "Reading guide":
                    paragraph.paragraph_format.page_break_before = True
            elif kind in ("paragraph", "text"):
                self.paragraph(token)
            elif kind == "list":
                self.list(token)
            elif kind == "table":
                self.table(token)
            elif kind == "code":
                self.code(token)
            elif kind == "blockquote":
                self.blocks(token["tokens"])
            elif kind == "hr":
                self.doc.add_paragraph("—" * 20)
            else:
                raise ValueError(f"Unsupported block token: {kind}")

    def export(self, path: Path):
        self.cover()
        self.blocks(self.tokens)
        expected = sum(1 for token in self.tokens if token["type"] == "code" and str(token.get("lang", "")).strip().startswith("mermaid"))
        if self.figure_count != expected:
            raise ValueError(f"Expected {expected} diagrams; embedded {self.figure_count}")
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        print(f"Exported {self.block_count} content blocks, {self.figure_count} diagrams, and {self.bookmark_count} heading bookmarks to {path}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--ast", required=True, type=Path)
    parser.add_argument("--figures", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    tokens = json.loads(args.ast.read_text(encoding="utf-8"))
    if isinstance(tokens, dict):
        tokens = tokens["tokens"]
    if not isinstance(tokens, list):
        parser.error("AST must be an array of Marked block tokens")
    Exporter(tokens, args.figures).export(args.output)


if __name__ == "__main__":
    main()
