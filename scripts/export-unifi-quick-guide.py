#!/usr/bin/env python3
"""Export the beginner-friendly Uni-Fi quick guide to DOCX."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]*\))", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)
        elif part.startswith("["):
            label, _, href = part[1:].partition("](")
            href = href[:-1] if href.endswith(")") else href
            run = paragraph.add_run(label)
            run.underline = True
            run.font.color.rgb = RGBColor(19, 115, 133)
        else:
            paragraph.add_run(part)


def page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def export(source: Path, target: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for level, size in ((1, 20), (2, 15), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(23, 62, 77)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("UNI-FI / MESHUSF  •  QUICK GUIDE").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("High-level project guide  |  Page ").font.size = Pt(8)
    page_field(footer)

    in_code = False
    code_lines: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.0
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Liberation Mono"
                run.font.size = Pt(8.5)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            doc.add_paragraph(heading.group(2), style=f"Heading {len(heading.group(1))}")
            i += 1
            continue
        if line.startswith("| ") and i + 1 < len(lines) and lines[i + 1].startswith("| ---"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [cell.strip() for cell in lines[i].strip("|").split("|")]
                if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    inline(table.cell(r, c).paragraphs[0], value)
                    if r == 0:
                        for run in table.cell(r, c).paragraphs[0].runs:
                            run.bold = True
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            inline(p, bullet.group(1))
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            inline(p, numbered.group(1))
            i += 1
            continue
        p = doc.add_paragraph()
        inline(p, line)
        i += 1
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Uni-Fi: quick guide"
    doc.core_properties.subject = "High-level MeshUSF project explanation"
    doc.core_properties.author = "MeshUSF"
    doc.save(target)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: export-unifi-quick-guide.py SOURCE.md OUTPUT.docx")
    export(Path(sys.argv[1]), Path(sys.argv[2]))
